"""Slide-preserving PPTX text extraction."""

from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path
from typing import TypedDict
from zipfile import BadZipFile

from pptx import Presentation
from pptx.exc import PackageNotFoundError


class ParsedSlide(TypedDict):
    """Text extracted from one 1-based PPTX slide."""

    page: int
    text: str


def parse_pptx(path: str | Path) -> list[ParsedSlide]:
    """Extract slide text while preserving 1-based slide numbers."""

    pptx_path = Path(path)
    if pptx_path.suffix.lower() != ".pptx":
        raise ValueError(f"Expected a PPTX file, got: {pptx_path}")
    if not pptx_path.is_file():
        raise FileNotFoundError(pptx_path)

    try:
        presentation = Presentation(pptx_path)
    except (BadZipFile, KeyError, PackageNotFoundError, ValueError, OSError) as exc:
        raise ValueError(f"Invalid PPTX file: {pptx_path}") from exc

    if len(presentation.slides) == 0:
        raise ValueError(f"PPTX contains zero slides: {pptx_path}")

    slides: list[ParsedSlide] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            parts.extend(_shape_text_parts(shape))
        parts.extend(_notes_text_parts(slide))
        slides.append({"page": slide_number, "text": "\n".join(parts)})
    return slides


def _shape_text_parts(shape: object) -> list[str]:
    parts: list[str] = []
    if getattr(shape, "has_text_frame", False):
        text = _text_frame_text(shape.text_frame)
        if text:
            parts.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            row_text = "\t".join(
                text
                for cell in row.cells
                if (text := _text_frame_text(cell.text_frame))
            )
            if row_text:
                parts.append(row_text)
    return parts


def _notes_text_parts(slide: object) -> list[str]:
    if not getattr(slide, "has_notes_slide", False):
        return []
    try:
        notes_text_frame = slide.notes_slide.notes_text_frame
    except (AttributeError, ValueError, KeyError):
        return []
    text = _text_frame_text(notes_text_frame)
    return [text] if text else []


def _text_frame_text(text_frame: object) -> str:
    lines = list(_paragraph_texts(text_frame.paragraphs))
    return "\n".join(lines).strip()


def _paragraph_texts(paragraphs: Iterable[object]) -> Iterable[str]:
    for paragraph in paragraphs:
        text = "".join(run.text for run in paragraph.runs).strip()
        if text:
            yield text
