from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from edurage.ingest import document_parser
from edurage.ingest.docx_parser import parse_docx


def _build_report(path: Path) -> Path:
    document = Document()
    document.add_heading("Chapter One", level=1)
    document.add_paragraph("First body paragraph.")
    document.add_paragraph("Second body paragraph.")
    document.add_heading("Chapter Two", level=2)
    document.add_paragraph("Third body paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "PC"
    table.cell(1, 1).text = "01H"
    document.save(path)
    return path


def test_parse_docx_splits_sections_on_heading_styles(tmp_path: Path) -> None:
    source = _build_report(tmp_path / "report.docx")

    pages = parse_docx(source)

    assert [page["page"] for page in pages] == [1, 2]
    assert pages[0]["text"] == "Chapter One\n\nFirst body paragraph.\n\nSecond body paragraph."
    assert pages[1]["text"].startswith("Chapter Two\n\nThird body paragraph.")
    assert pages[1]["text"].splitlines()[0] == "Chapter Two"
    assert all(page["text_sources"] == ["text"] for page in pages)


def test_parse_docx_renders_table_rows_with_tabs(tmp_path: Path) -> None:
    source = _build_report(tmp_path / "report.docx")

    pages = parse_docx(source)

    table_lines = pages[1]["text"].splitlines()
    assert "表格 2行×2列" in table_lines
    assert "Name\tValue" in table_lines
    assert "PC\t01H" in table_lines


def test_parse_docx_without_headings_breaks_on_blank_paragraphs(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("")
    document.add_paragraph("Second paragraph.")
    source = tmp_path / "plain.docx"
    document.save(source)

    pages = parse_docx(source)

    assert [page["text"] for page in pages] == ["First paragraph.", "Second paragraph."]
    assert [page["page"] for page in pages] == [1, 2]


def test_parse_docx_rejects_wrong_suffix(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="Expected a DOCX file"):
        parse_docx(tmp_path / "report.pdf")


def test_parse_docx_rejects_invalid_package(tmp_path: Path) -> None:
    source = tmp_path / "broken.docx"
    source.write_bytes(b"PK\x03\x04this is not a docx package")

    with pytest.raises(ValueError, match="Invalid DOCX file"):
        parse_docx(source)


def test_parse_docx_requires_existing_file(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        parse_docx(tmp_path / "missing.docx")


def test_parse_docx_rejects_document_without_text(tmp_path: Path) -> None:
    source = tmp_path / "empty.docx"
    Document().save(source)

    with pytest.raises(ValueError, match="no extractable text"):
        parse_docx(source)


def test_document_parser_dispatches_real_docx_file(tmp_path: Path) -> None:
    source = _build_report(tmp_path / "report.docx")

    assert document_parser.detect_file_type(source) == "docx"
    pages = document_parser.parse_document(source)

    assert [page["source_type"] for page in pages] == ["docx", "docx"]
    assert all(page["text_sources"] == ["text"] for page in pages)
    assert all(page["page_used_ocr"] is False for page in pages)
