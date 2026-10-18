from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient

from edurage.api.deps import get_app_settings, get_indexer
from edurage.core.config import Settings
from edurage.ingest.embedder import HashingEmbedder
from edurage.ingest.indexer import IngestIndexer
from tests.test_ingest_pipeline import FakeRepository, FakeVectorStore

DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


class FakeIndexer:
    """Capture the async upload handshake without parsing the document."""

    def __init__(self) -> None:
        self.prepared: list[dict[str, Any]] = []
        self.ingested: list[dict[str, Any]] = []

    def prepare_document(
        self,
        path: Path,
        course_code: str,
        *,
        original_filename: str | None = None,
        file_type: str | None = None,
    ) -> str:
        self.prepared.append(
            {
                "path": path,
                "course_code": course_code,
                "original_filename": original_filename,
                "file_type": file_type,
            }
        )
        return "docx-123"

    def ingest_document(
        self,
        path: Path,
        course_code: str,
        *,
        original_filename: str | None = None,
        document_id: str | None = None,
    ) -> dict[str, Any]:
        self.ingested.append(
            {
                "path": path,
                "course_code": course_code,
                "original_filename": original_filename,
                "document_id": document_id,
            }
        )
        return {"document_id": document_id, "pages": 1, "chunks": 1}


class FakeObjectStore:
    def __init__(self) -> None:
        self.objects: dict[str, tuple[bytes, str]] = {}

    def put(self, key: str, body: bytes, content_type: str) -> None:
        self.objects[key] = (body, content_type)

    def delete(self, key: str) -> None:
        self.objects.pop(key, None)


def _docx_bytes(tmp_path: Path) -> bytes:
    document = Document()
    document.add_heading("Chapter One", level=1)
    document.add_paragraph("Computer organization body text with enough words to chunk.")
    source = tmp_path / "build.docx"
    document.save(source)
    return source.read_bytes()


def _client_with_upload_overrides(
    client: TestClient, fake_indexer: FakeIndexer
) -> None:
    client.app.dependency_overrides[get_indexer] = lambda: fake_indexer
    client.app.dependency_overrides[get_app_settings] = lambda: Settings(
        _env_file=None, allow_cloud_upload=True
    )


def test_docx_upload_is_accepted_with_docx_content_type(
    client: TestClient, tmp_path: Path
) -> None:
    fake_indexer = FakeIndexer()
    _client_with_upload_overrides(client, fake_indexer)

    response = client.post(
        "/api/v1/upload",
        files={"file": ("课程讲义.docx", _docx_bytes(tmp_path), DOCX_CONTENT_TYPE)},
    )

    assert response.status_code == 202
    assert response.json() == {
        "document_id": "docx-123",
        "filename": "课程讲义.docx",
        "course_code": "CS-COMP-ORG",
        "status": "processing",
    }
    assert fake_indexer.prepared[0]["file_type"] == "docx"
    assert fake_indexer.ingested[0]["document_id"] == "docx-123"


def test_docx_upload_rejects_mime_mismatch(
    client: TestClient, tmp_path: Path
) -> None:
    fake_indexer = FakeIndexer()
    _client_with_upload_overrides(client, fake_indexer)

    response = client.post(
        "/api/v1/upload",
        files={"file": ("课程讲义.docx", _docx_bytes(tmp_path), "application/pdf")},
    )

    assert response.status_code == 400
    assert response.json()["error"]["code"] == "mime_mismatch"
    assert fake_indexer.prepared == []


def test_docx_upload_rejects_zip_without_document_part(client: TestClient) -> None:
    fake_indexer = FakeIndexer()
    _client_with_upload_overrides(client, fake_indexer)
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("word/other.xml", "<xml />")

    response = client.post(
        "/api/v1/upload",
        files={"file": ("broken.docx", buffer.getvalue(), DOCX_CONTENT_TYPE)},
    )

    assert response.status_code == 400
    assert response.json()["error"]["code"] == "invalid_upload"
    assert fake_indexer.prepared == []


def test_docx_ingest_persists_sections_chunks_and_object(tmp_path: Path) -> None:
    source = tmp_path / "lesson.docx"
    source.write_bytes(_docx_bytes(tmp_path))
    settings = Settings(_env_file=None, ocr_enabled=False)
    repository = FakeRepository()
    vectors = FakeVectorStore()
    objects = FakeObjectStore()
    indexer = IngestIndexer(
        embedder=HashingEmbedder(),
        vector_store=vectors,
        repository=repository,
        object_store=objects,
        settings=settings,
    )

    document_id = indexer.prepare_document(
        source, "cs-demo", original_filename="lesson.docx", file_type="docx"
    )
    assert repository.documents[document_id]["status"] == "processing"
    assert repository.documents[document_id]["metadata"]["progress"]["stage"] == "parse"

    result = indexer.ingest_document(
        source, "cs-demo", original_filename="lesson.docx", document_id=document_id
    )

    assert result["file_type"] == "docx"
    assert result["pages"] >= 1
    assert repository.documents[document_id]["status"] == "indexed"
    assert vectors.points
    assert {point["payload"]["file_type"] for point in vectors.points} == {"docx"}
    assert {point["payload"]["source"] for point in vectors.points} == {"docx"}
    (object_key, (body, content_type)), = objects.objects.items()
    assert content_type == DOCX_CONTENT_TYPE
    assert body == source.read_bytes()
    assert object_key.endswith("/lesson.docx")
    assert source.exists()
