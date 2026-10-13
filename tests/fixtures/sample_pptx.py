"""Deterministic three-slide PPTX fixture."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


def make_sample_pptx(path: Path) -> Path:
    """Write a deterministic three-slide PPTX and return its path."""

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()

    _add_text_slide(
        presentation,
        "EduRAG sample PPTX slide 1",
        "Course Pack: Computer Organization. This slide introduces the deck fixture.",
    )
    _add_table_slide(presentation)
    _add_text_slide(
        presentation,
        "EduRAG sample PPTX slide 3",
        "Chunk metadata includes the source slide number. Stable text closes the deck.",
    )

    presentation.save(path)
    return path


def _add_text_slide(presentation: Presentation, title: str, body: str) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = title
    text_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.4), Inches(1.2))
    text_box.text_frame.text = body
    empty_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(2.0), Inches(0.5))
    empty_box.text_frame.text = "   "


def _add_table_slide(presentation: Presentation) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "EduRAG sample PPTX slide 2"
    body = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.4), Inches(1.0))
    body.text_frame.text = "ADD 指令的微地址序列是取指、译码、执行。"
    table_shape = slide.shapes.add_table(2, 2, Inches(0.8), Inches(2.6), Inches(6.0), Inches(1.2))
    table = table_shape.table
    table.cell(0, 0).text = "Phase"
    table.cell(0, 1).text = "Description"
    table.cell(1, 0).text = "Execute"
    table.cell(1, 1).text = "Uses slide-aware ingestion."
