"""Offline regressions for ingestion quality, provenance, and filename contracts."""

from __future__ import annotations

import inspect
import sys
from collections.abc import Mapping, Sequence
from itertools import pairwise
from pathlib import Path
from typing import Any

import pymupdf
import pytest
from pptx import Presentation
from pptx.util import Inches

from edurage.core.config import Settings
from edurage.ingest import document_parser
from edurage.ingest.embedder import HashingEmbedder
from edurage.ingest.indexer import IngestIndexer
from edurage.ingest.ocr import OcrPageResult
from edurage.ingest.pdf_parser import parse_pdf
from edurage.ingest.pptx_parser import parse_pptx
from edurage.ingest.structured_chunker import structured_chunk_pages
from tests.fixtures.sample_pdf import make_sample_pdf
from tests.fixtures.sample_pptx import make_sample_pptx
from tests.test_pptx_ingest import FakeObjectStore, FakeRepository, FakeVectorStore


@pytest.mark.parametrize("overlap", [0, 1, 19, 80, 499])
@pytest.mark.parametrize("length", [500, 501, 840, 841, 920, 1000, 1421])
def test_long_text_windows_reconstruct_input_without_redundant_tail(
    length: int, overlap: int
) -> None:
    text = "".join(str(number % 10) for number in range(length))
    chunks = structured_chunk_pages([{"page": 1, "text": text}], overlap=overlap)

    reconstructed = chunks[0]["text"]
    for previous, current in pairwise(chunks):
        if overlap:
            assert previous["text"][-overlap:] == current["text"][:overlap]
        assert len(current["text"]) > overlap
        reconstructed += current["text"][overlap:]
    assert reconstructed == text
    assert all(0 < len(chunk["text"]) <= 500 for chunk in chunks)
    assert min(len(chunk["text"]) for chunk in chunks) >= 20
    assert [chunk["chunk_index"] for chunk in chunks] == list(range(len(chunks)))


def test_many_short_cells_and_paragraphs_keep_every_value_in_order() -> None:
    cells = ["A", "1", "B", "2", "PC", "00H", "IR", "01H", "0", "1"]
    text = "\n\n".join(cells * 12)
    chunks = structured_chunk_pages([{"page": 13, "text": text}], overlap=0)
    assert len(chunks) <= 3
    assert "\n".join(chunk["text"] for chunk in chunks).splitlines() == cells * 12
    assert all(chunk["page"] == 13 for chunk in chunks)
    assert all(chunk["section"] is None for chunk in chunks)


def test_independent_sections_keep_their_own_title_and_body() -> None:
    chunks = structured_chunk_pages(
        [{"page": 1, "text": "# First\n\nFirst body.\n\n# Second\n\nSecond body."}]
    )
    assert [chunk["section"] for chunk in chunks] == ["# First", "# Second"]
    assert "First body." in chunks[0]["text"]
    assert "Second body." in chunks[1]["text"]
    assert all(chunk["block_type"] == "paragraph" for chunk in chunks)


def test_short_caption_attaches_to_following_description() -> None:
    chunks = structured_chunk_pages(
        [{"page": 1, "text": "Figure 2 Circuit\n\nSignals follow the numbered connections."}]
    )
    assert len(chunks) == 1
    assert chunks[0]["block_type"] == "caption"
    assert chunks[0]["text"].splitlines() == [
        "Figure 2 Circuit", "Signals follow the numbered connections."
    ]


def test_structure_metadata_only_describes_content_inside_each_window() -> None:
    chunks = structured_chunk_pages(
        [{"page": 1, "text": "Figure 2 Circuit\n\n" + "a" * 200}],
        chunk_size=80,
        overlap=20,
    )
    assert chunks[0]["block_type"] == "caption"
    assert all(chunk["block_type"] == "paragraph" for chunk in chunks[1:])


def test_provenance_is_calculated_for_each_window_including_overlap() -> None:
    chunks = structured_chunk_pages(
        [
            {
                "page": 1,
                "text": "ID\n\n[OCR]\n" + "x" * 160,
                "text_segments": [
                    {"text": "ID", "source": "text"},
                    {"text": "x" * 160, "source": "ocr"},
                ],
                "text_sources": ["text", "ocr"],
                "page_used_ocr": True,
                "ocr_confidence": 0.75,
                "source": "pdf",
                "section": "Evidence",
            }
        ],
        chunk_size=80,
        overlap=20,
    )
    assert [chunk["modality"] for chunk in chunks] == ["mixed", "ocr", "ocr"]
    assert chunks[0]["text_sources"] == ["text", "ocr"]
    assert all(chunk["text_sources"] == ["ocr"] for chunk in chunks[1:])
    assert all(chunk["page_used_ocr"] for chunk in chunks)
    assert all(chunk["ocr_confidence"] == 0.75 for chunk in chunks)
    assert all(chunk["source_type"] == "pdf" for chunk in chunks)
    assert all(chunk["section"] == "Evidence" for chunk in chunks)
    assert all("[OCR]" not in chunk["text"] for chunk in chunks)


def test_text_blocks_remain_text_on_a_page_with_ocr() -> None:
    chunks = structured_chunk_pages(
        [
            {
                "page": 1,
                "text": "t" * 120 + "\n[OCR]\n" + "o" * 120,
                "text_sources": ["text", "ocr"],
                "page_used_ocr": True,
            }
        ]
    )
    assert [chunk["modality"] for chunk in chunks] == ["text", "ocr"]
    assert [chunk["text_sources"] for chunk in chunks] == [["text"], ["ocr"]]
    assert all(chunk["block_type"] == "paragraph" for chunk in chunks)
    assert all(chunk["page_used_ocr"] for chunk in chunks)


def test_ambiguous_legacy_sources_are_mixed_and_failure_metadata_is_preserved() -> None:
    chunk = structured_chunk_pages(
        [
            {
                "page": 1,
                "text": "Legacy merged text",
                "text_sources": ["text", "ocr"],
                "ocr_error": "original diagnostic",
            }
        ]
    )[0]
    assert chunk["modality"] == "mixed"
    assert chunk["text_sources"] == ["text", "ocr"]
    assert chunk["page_used_ocr"] is True
    assert chunk["ocr_error"] == "original diagnostic"


def test_legacy_repeated_ocr_markers_do_not_become_content() -> None:
    chunks = structured_chunk_pages(
        [{"page": 1, "text": "native\n[OCR]\nimage one\n[OCR]\nimage two"}]
    )
    assert len(chunks) == 1
    assert chunks[0]["modality"] == "mixed"
    assert chunks[0]["text"] == "native\nimage one\nimage two"


class FakeOCR:
    def __init__(self, text: str) -> None:
        self.text = text
        self.calls: list[tuple[int, str]] = []

    def recognize(self, image: bytes, *, page: int, language: str) -> OcrPageResult:
        assert image.startswith(b"\x89PNG")
        self.calls.append((page, language))
        return OcrPageResult(page=page, text=self.text, confidence=0.9)


def test_duplicate_and_empty_ocr_results_still_record_execution(tmp_path: Path) -> None:
    source = make_sample_pdf(tmp_path / "sample.pdf")
    original = parse_pdf(source)[0]["text"]
    for recognized in (original, ""):
        provider = FakeOCR(recognized)
        page = parse_pdf(
            source, ocr_provider=provider, ocr_enabled=True, ocr_min_text_chars=10000
        )[0]
        assert page["page_used_ocr"] is True
        assert page["text"] == original
        assert page["text_sources"] == ["text"]
        assert [segment["source"] for segment in page["text_segments"]] == ["text"]
        chunks = structured_chunk_pages([page])
        assert all(chunk["modality"] == "text" for chunk in chunks)
        assert all(chunk["page_used_ocr"] for chunk in chunks)


@pytest.mark.parametrize("entrypoint", ["ingest_document", "ingest_pdf", "ingest_pptx"])
@pytest.mark.parametrize(
    ("original", "expected"),
    [
        (None, "temporary-upload"),
        ("/client/folder/lesson.pdf", "lesson.pdf"),
        ("C:\\fakepath\\Lesson.pptx", "Lesson.pptx"),
        ("../  lesson:\x00?.pdf.  ", "lesson__.pdf"),
    ],
)
def test_filename_contract_across_all_entrypoints(
    tmp_path: Path, entrypoint: str, original: str | None, expected: str
) -> None:
    pptx = entrypoint == "ingest_pptx"
    suffix = ".pptx" if pptx else ".pdf"
    path = tmp_path / ("temporary-upload" + suffix)
    (make_sample_pptx if pptx else make_sample_pdf)(path)
    vectors, repository, objects = FakeVectorStore(), FakeRepository(), FakeObjectStore()
    indexer = IngestIndexer(
        embedder=HashingEmbedder(), vector_store=vectors, repository=repository,
        object_store=objects, settings=Settings(_env_file=None, ocr_enabled=False),
    )
    method = getattr(indexer, entrypoint)
    assert inspect.signature(method).parameters["original_filename"].kind == (
        inspect.Parameter.KEYWORD_ONLY
    )
    result = method(path, "TEST", original_filename=original)
    expected_name = expected + suffix if original is None else expected
    document = repository.documents[result["document_id"]]
    assert document["filename"] == expected_name
    assert document["storage_path"].split("/")[-1] == expected_name
    assert set(objects.objects) == {document["storage_path"]}
    assert all(point["payload"]["filename"] == expected_name for point in vectors.points)
    assert all(chunk["metadata"]["filename"] == expected_name for chunk in repository.chunks)


@pytest.mark.parametrize("failure", ["ensure", "upsert", "document", "chunks", "status"])
def test_original_filename_compensation_cleans_partial_writes(
    tmp_path: Path, failure: str
) -> None:
    class Vectors(FakeVectorStore):
        def ensure_collection(self, collection_name: str, vector_size: int) -> None:
            super().ensure_collection(collection_name, vector_size)
            if failure == "ensure":
                raise RuntimeError("injected")

        def upsert(self, collection_name: str, points: Sequence[Mapping[str, Any]]) -> None:
            self.points.append(points[0])
            if failure == "upsert":
                raise RuntimeError("injected")
            self.points.extend(points[1:])

    class Repository(FakeRepository):
        def insert_document(self, document: Mapping[str, Any]) -> None:
            super().insert_document(document)
            if failure == "document":
                raise RuntimeError("injected")

        def insert_chunks(self, chunks: Sequence[Mapping[str, Any]]) -> None:
            self.chunks.append(chunks[0])
            if failure == "chunks":
                raise RuntimeError("injected")
            self.chunks.extend(chunks[1:])

        def set_document_status(self, document_id: str, status: str) -> None:
            if failure == "status":
                raise RuntimeError("injected")
            super().set_document_status(document_id, status)

    source = make_sample_pdf(tmp_path / "temporary.pdf")
    vectors, repository, objects = Vectors(), Repository(), FakeObjectStore()
    indexer = IngestIndexer(
        embedder=HashingEmbedder(), vector_store=vectors, repository=repository,
        object_store=objects, settings=Settings(_env_file=None, ocr_enabled=False),
    )
    with pytest.raises(RuntimeError, match="injected"):
        indexer.ingest_document(source, "TEST", original_filename="../original.pdf")
    assert not vectors.points
    assert not repository.documents
    assert not repository.chunks
    assert not objects.objects


def test_fake_ocr_provenance_reaches_both_stores(tmp_path: Path) -> None:
    path = tmp_path / "upload.pdf"
    with pymupdf.open() as document:
        page = document.new_page()
        for row in range(4):
            page.insert_text((72, 72 + row * 20), "Embedded source words repeated here.")
        document.save(path)
    vectors, repository, objects = FakeVectorStore(), FakeRepository(), FakeObjectStore()
    provider = FakeOCR("Recognized image words. " * 40)
    indexer = IngestIndexer(
        embedder=HashingEmbedder(), vector_store=vectors, repository=repository,
        object_store=objects, ocr_provider=provider,
        settings=Settings(_env_file=None, ocr_enabled=True, ocr_min_text_chars=1000),
    )
    indexer.ingest_document(path, "TEST", original_filename="original.pdf")
    assert provider.calls == [(1, "chi_sim+eng")]
    assert {point["payload"]["modality"] for point in vectors.points} == {"text", "ocr"}
    for point, chunk in zip(vectors.points, repository.chunks, strict=True):
        assert point["payload"]["modality"] == chunk["modality"]
        for field in (
            "filename", "page_used_ocr", "text_sources", "ocr_confidence", "ocr_error",
            "source", "section", "block_type",
        ):
            assert point["payload"][field] == chunk["metadata"][field]
        assert chunk["metadata"]["page_used_ocr"] is True
        assert chunk["metadata"]["ocr_confidence"] == 0.9
        assert "[OCR]" not in chunk["content"]


@pytest.mark.parametrize("cli", ["document", "pdf"])
def test_both_dry_run_clis_use_settings_provider_and_structured_chunks(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str],
    cli: str,
) -> None:
    from scripts import ingest_local_document, ingest_local_pdf

    source = make_sample_pdf(tmp_path / "sample.pdf")
    settings = Settings(
        _env_file=None, ocr_enabled=True, ocr_language="eng",
        ocr_min_text_chars=10000, chunk_size=120, chunk_overlap=20,
        ocr_tesseract_cmd="fake-tesseract",
    )
    provider = FakeOCR("a\n1\nb\n2\nc\n3")
    builds: list[tuple[bool, str, str]] = []
    observed: list[dict[str, Any]] = []

    def build(enabled: bool, engine: str, *, tesseract_cmd: str) -> FakeOCR:
        builds.append((enabled, engine, tesseract_cmd))
        return provider

    def chunk(pages: Any, **kwargs: Any) -> Any:
        result = structured_chunk_pages(pages, **kwargs)
        observed.extend(result)
        return result

    monkeypatch.setattr(document_parser, "get_settings", lambda: settings)
    monkeypatch.setattr(document_parser, "build_ocr_provider", build)
    monkeypatch.setattr(ingest_local_document, "get_settings", lambda: settings)
    monkeypatch.setattr(ingest_local_pdf, "get_settings", lambda: settings)
    monkeypatch.setattr(ingest_local_document, "structured_chunk_pages", chunk)
    monkeypatch.setattr(
        sys, "argv",
        ["ingest", "--" + cli, str(source), "--course-code", "TEST", "--dry-run"],
    )
    module = ingest_local_document if cli == "document" else ingest_local_pdf
    assert module.main() == 0
    assert builds == [(True, "tesseract", "fake-tesseract")]
    assert provider.calls == [(1, "eng"), (2, "eng"), (3, "eng")]
    assert observed
    assert all(len(item["text"]) <= 120 for item in observed)
    assert any(item["modality"] in {"ocr", "mixed"} for item in observed)
    assert "embedding=hashing" in capsys.readouterr().out


def test_pptx_table_retains_empty_columns(tmp_path: Path) -> None:
    path = tmp_path / "cells.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(
        2, 3, Inches(1), Inches(1), Inches(6), Inches(2)
    ).table
    for row, values in zip(table.rows, [("", "PC", "01H"), ("IR", "", "02H")], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    presentation.save(path)
    pages = parse_pptx(path)
    assert pages[0]["text"].splitlines() == ["\tPC\t01H", "IR\t\t02H"]
    chunk = structured_chunk_pages(pages)[0]
    assert chunk["block_type"] == "table"
    assert chunk["text"].splitlines() == ["\tPC\t01H", "IR\t\t02H"]
